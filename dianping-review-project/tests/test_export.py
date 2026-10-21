import csv
import json
from openpyxl import load_workbook
from processing.models import Review
from export.export_excel import export_reviews, FIELDS


def test_exports_evidence_quality_and_formula_safety(tmp_path):
    good=Review('good',4.5,'6月28日','','=这不是公式',['a.png','b.png'],.95,'needs_review','time',['year_missing'])
    fragment=Review('fragment',None,'','','片段不计入评论数量',['c.png'],.5,'incomplete','time',
                    ['unjoined_fragment'],fragment=True)
    excluded=Review('duplicate',4.5,'6月28日','','重复记录',['d.png'],.95,'needs_review','time',excluded=True)
    result=export_reviews([good,fragment,excluded],tmp_path)
    assert result['records']==1 and result['fragments']==1 and result['quality_rows']==3
    csv_path,xlsx_path,quality_path=result['paths']
    with open(csv_path,encoding='utf-8-sig',newline='') as file:
        reader=csv.DictReader(file)
        rows=list(reader)
        assert reader.fieldnames==FIELDS
    assert rows[0]['review_text']=="'=这不是公式"
    assert json.loads(rows[0]['source_images'])==['a.png','b.png']
    workbook=load_workbook(xlsx_path)
    assert workbook.active.cell(2,5).value=='=这不是公式'
    assert workbook.active.cell(2,5).data_type=='s'
    assert load_workbook(quality_path).active.max_row==4
    assert good.review_text=='=这不是公式'


def test_empty_export_has_headers(tmp_path):
    result=export_reviews([],tmp_path,demo=True)
    assert result['records']==0
    assert all('演示_非真实数据' in path for path in result['paths'])
    assert load_workbook(result['paths'][1]).active.max_row==1
